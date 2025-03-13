import os
import sys
from docx import Document

def resource_path(relative_path):
    """Obter o caminho absoluto para o recurso, funciona para dev e PyInstaller"""
    try:
        # PyInstaller cria uma pasta temporária e armazena o caminho nela
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.abspath(".")

    return os.path.join(base_path, relative_path)

def get_user_documents_path():
    """Obter o caminho para a pasta de documentos do usuário"""
    return os.path.join(os.path.expanduser("~"), "Documents")

def generate_contract(dados_contrato):
    template_path = resource_path(os.path.join('templates', 'contrato_modelo.docx'))
    doc = Document(template_path)

    # Substituir os placeholders no documento com os dados do contrato
    for p in doc.paragraphs:
        for key, value in dados_contrato.items():
            if f'{{{{{key}}}}}' in p.text:
                p.text = p.text.replace(f'{{{{{key}}}}}', str(value))

    # Salvar o documento gerado na pasta de documentos do usuário
    output_path = os.path.join(get_user_documents_path(), f'contract_generated{dados_contrato["contratante"]}.docx')
    doc.save(output_path)
    return output_path